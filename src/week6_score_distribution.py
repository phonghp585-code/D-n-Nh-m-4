"""Descriptive comparison of Vietnamese THPT exam score distributions (2023 vs 2025).

Uses the already-harmonized datasets from branches thanh (2023) and nhan (2025).
No imputation is performed. Scores outside [0, 10] are excluded per subject.
"""
from __future__ import annotations
import argparse, json, re, unicodedata
from pathlib import Path
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

SUBJECT_ALIASES = {
    "Toán": ["mathematics", "math", "toan", "điểm toán", "toán"],
    "Ngữ văn": ["literature", "ngu van", "ngữ văn", "van", "điểm văn"],
    "Ngoại ngữ": ["foreign language", "foreign_language", "ngoai ngu", "ngoại ngữ", "english"],
    "Vật lý": ["physics", "vat ly", "vật lý", "ly"],
    "Hóa học": ["chemistry", "hoa hoc", "hóa học", "hoa"],
    "Sinh học": ["biology", "sinh hoc", "sinh học", "sinh"],
    "Lịch sử": ["history", "lich su", "lịch sử", "su"],
    "Địa lý": ["geography", "dia ly", "địa lý", "dia"],
    "GDCD": ["civic education", "gdcd", "giao duc cong dan", "giáo dục công dân"],
    "GDKT&PL": ["economic and legal education", "gdkt&pl", "giáo dục kinh tế và pháp luật"],
}

def norm(value):
    value = unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "", value)

def read_dataset(path):
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Không tìm thấy dataset: {path}")
    # Detect an unhydrated Git LFS pointer instead of reporting confusing CSV errors.
    with path.open("rb") as f:
        prefix = f.read(100)
    if prefix.startswith(b"version https://git-lfs.github.com/spec/v1"):
        raise RuntimeError(f"{path} vẫn là Git LFS pointer. Hãy chạy 'git lfs pull' để tải dữ liệu thật.")
    if path.suffix.lower() == ".parquet":
        return pd.read_parquet(path)
    if path.suffix.lower() in (".xlsx", ".xls"):
        return pd.read_excel(path)
    return pd.read_csv(path, low_memory=False)

def map_subject_columns(df):
    normalized = {norm(c): c for c in df.columns}
    found = {}
    for subject, aliases in SUBJECT_ALIASES.items():
        for alias in aliases:
            key = norm(alias)
            if key in normalized:
                found[subject] = normalized[key]
                break
        if subject not in found:
            for key, original in normalized.items():
                if any(len(norm(a)) >= 6 and (norm(a) in key or key in norm(a)) for a in aliases):
                    found[subject] = original
                    break
    return found

def tidy_year(df, year):
    columns = map_subject_columns(df)
    if not columns:
        raise ValueError(f"Năm {year}: không nhận diện được cột môn. Cột hiện có: {list(df.columns)}")
    result = pd.DataFrame({"year": year}, index=df.index)
    for subject, column in columns.items():
        values = pd.to_numeric(df[column].replace({"": np.nan, "-": np.nan}), errors="coerce")
        result[subject] = values.where(values.between(0, 10))
    return result.reset_index(drop=True)

def summarize(long):
    rows = []
    for (year, subject), group in long.groupby(["year", "subject"], dropna=False):
        scores = group["score"].dropna()
        if scores.empty:
            continue
        rows.append({
            "year": int(year), "subject": subject, "n": int(scores.size),
            "mean": scores.mean(), "median": scores.median(), "std": scores.std(ddof=1),
            "skewness": scores.skew(), "kurtosis_excess": scores.kurt(),
            "min": scores.min(), "q1": scores.quantile(.25), "q3": scores.quantile(.75), "max": scores.max(),
            "pct_le_1": (scores <= 1).mean() * 100, "pct_lt_5": (scores < 5).mean() * 100,
            "pct_ge_8": (scores >= 8).mean() * 100, "pct_ge_9": (scores >= 9).mean() * 100,
        })
    return pd.DataFrame(rows).sort_values(["subject", "year"]).reset_index(drop=True)

def slug(text):
    return re.sub(r"[^a-z0-9]+", "_", norm(text)).strip("_")

def plot_subject(long, subject, output_dir):
    subset = long[(long.subject == subject) & long.score.notna()]
    if subset.empty:
        return []
    output_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    years = sorted(subset.year.unique())
    # Histogram with common bins and density scale.
    fig, ax = plt.subplots(figsize=(9, 5))
    bins = np.arange(0, 10.25, .25)
    for year in years:
        values = subset.loc[subset.year == year, "score"]
        ax.hist(values, bins=bins, density=True, alpha=.4, label=str(year))
    ax.set(title=f"Phân phối điểm {subject}: histogram 2023–2025", xlabel="Điểm", ylabel="Mật độ")
    ax.set_xlim(0, 10); ax.legend(); ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    path = output_dir / f"{slug(subject)}_histogram.png"
    fig.savefig(path, dpi=180, bbox_inches="tight"); plt.close(fig); paths.append(path)

    # KDE is smoothed; fallback to step-density histogram if scipy/KDE is unavailable or singular.
    fig, ax = plt.subplots(figsize=(9, 5))
    for year in years:
        values = subset.loc[subset.year == year, "score"].to_numpy()
        try:
            from scipy.stats import gaussian_kde
            grid = np.linspace(0, 10, 400)
            ax.plot(grid, gaussian_kde(values)(grid), label=str(year))
        except Exception:
            ax.hist(values, bins=bins, density=True, histtype="step", linewidth=1.6, label=f"{year} (mật độ histogram)")
    ax.set(title=f"Đường mật độ điểm {subject}: 2023–2025", xlabel="Điểm", ylabel="Mật độ")
    ax.set_xlim(0, 10); ax.legend(); ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    path = output_dir / f"{slug(subject)}_kde.png"
    fig.savefig(path, dpi=180, bbox_inches="tight"); plt.close(fig); paths.append(path)

    fig, ax = plt.subplots(figsize=(8, 5))
    data = [subset.loc[subset.year == year, "score"].dropna() for year in years]
    ax.boxplot(data, labels=[str(year) for year in years], showfliers=False)
    ax.set(title=f"Boxplot điểm {subject}: 2023–2025", xlabel="Năm", ylabel="Điểm", ylim=(0, 10))
    ax.spines[["top", "right"]].set_visible(False); ax.grid(axis="y", alpha=.25)
    fig.tight_layout()
    path = output_dir / f"{slug(subject)}_boxplot.png"
    fig.savefig(path, dpi=180, bbox_inches="tight"); plt.close(fig); paths.append(path)

    fig, ax = plt.subplots(figsize=(8, 5))
    ax.violinplot([subset.loc[subset.year == year, "score"].dropna().to_numpy() for year in years],
                  showmedians=True, showextrema=True)
    ax.set_xticks(range(1, len(years) + 1), [str(year) for year in years])
    ax.set(title=f"Violin plot điểm {subject}: 2023–2025", xlabel="Năm", ylabel="Điểm", ylim=(0, 10))
    ax.spines[["top", "right"]].set_visible(False); ax.grid(axis="y", alpha=.25)
    fig.tight_layout()
    path = output_dir / f"{slug(subject)}_violin.png"
    fig.savefig(path, dpi=180, bbox_inches="tight"); plt.close(fig); paths.append(path)
    return paths

def write_docx(stats, figures, output_path, row_count, source23, source25):
    from docx import Document
    from docx.shared import Inches
    doc = Document()
    doc.add_heading("BÁO CÁO PHÂN TÍCH VÀ SO SÁNH PHỔ ĐIỂM THPT 2023–2025", 0)
    doc.add_paragraph("Dự án KHDL – Nhóm 4 | Tuần 6: thống kê mô tả; Tuần 7: trực quan hóa dữ liệu")
    doc.add_heading("1. Mục tiêu và dữ liệu", level=1)
    doc.add_paragraph("Mục tiêu: mô tả và so sánh phân phối điểm từng môn giữa năm 2023 và 2025 bằng thống kê vị trí, độ phân tán, hình dạng và tỷ lệ theo ngưỡng điểm.")
    doc.add_paragraph(f"Nguồn 2023: {source23}. Nguồn 2025: {source25}. Số quan sát điểm hợp lệ ở dạng dài: {row_count:,}.")
    doc.add_paragraph("Các môn được so sánh là giao của những môn nhận diện được ở cả hai dataset. Điểm thiếu không được điền; giá trị ngoài [0, 10] không được dùng trong thống kê.")
    doc.add_heading("2. Thống kê mô tả", level=1)
    table = doc.add_table(rows=1, cols=8); table.style = "Light Shading Accent 1"
    headers = ["Năm", "Môn", "N", "Trung bình", "Trung vị", "ĐLC", "Skewness", "Kurtosis excess"]
    for cell, value in zip(table.rows[0].cells, headers): cell.text = value
    for _, row in stats.iterrows():
        cells = table.add_row().cells
        values = [int(row.year), row.subject, int(row.n), f"{row['mean']:.3f}", f"{row['median']:.3f}",
                  f"{row['std']:.3f}", f"{row['skewness']:.3f}", f"{row['kurtosis_excess']:.3f}"]
        for cell, value in zip(cells, values): cell.text = str(value)
    doc.add_paragraph("Skewness > 0 biểu thị đuôi phải dài hơn; < 0 biểu thị lệch trái; gần 0 gần đối xứng. Kurtosis trong bảng là excess kurtosis, với phân phối chuẩn có giá trị tham chiếu 0. Các chỉ số hình dạng nhạy với cỡ mẫu và cách chấm điểm.")
    doc.add_heading("3. Tỷ lệ theo ngưỡng điểm", level=1)
    table = doc.add_table(rows=1, cols=6); table.style = "Light Shading Accent 1"
    for cell, value in zip(table.rows[0].cells, ["Năm", "Môn", "≤1 (%)", "<5 (%)", "≥8 (%)", "≥9 (%)"]): cell.text = value
    for _, row in stats.iterrows():
        cells = table.add_row().cells
        values = [int(row.year), row.subject, f"{row.pct_le_1:.2f}", f"{row.pct_lt_5:.2f}", f"{row.pct_ge_8:.2f}", f"{row.pct_ge_9:.2f}"]
        for cell, value in zip(cells, values): cell.text = str(value)
    doc.add_paragraph("Mỗi tỷ lệ dùng số điểm hợp lệ của chính môn và năm đó làm mẫu số. Vì vậy, cần đọc kèm cột N để nhận biết chênh lệch về số quan sát.")
    doc.add_heading("3.1. Bảng chênh lệch 2025 so với 2023", level=2)
    doc.add_paragraph("Các chênh lệch dưới đây được tính bằng giá trị năm 2025 trừ giá trị năm 2023. Với tỷ lệ, đơn vị là điểm phần trăm; với thống kê điểm, đơn vị là điểm trừ khi chỉ số không có đơn vị.")
    table = doc.add_table(rows=1, cols=10); table.style = "Light Shading Accent 1"
    delta_headers = ["Môn", "Δ TB", "Δ Trung vị", "Δ ĐLC", "Δ Skewness", "Δ Kurtosis", "Δ ≤1 (đpt)", "Δ <5 (đpt)", "Δ ≥8 (đpt)", "Δ ≥9 (đpt)"]
    for cell, value in zip(table.rows[0].cells, delta_headers): cell.text = value
    for subject in sorted(stats["subject"].unique()):
        part = stats[stats["subject"] == subject].set_index("year")
        if 2023 not in part.index or 2025 not in part.index:
            cells = table.add_row().cells; cells[0].text = subject
            for cell in cells[1:]: cell.text = "N/A"
            continue
        a, b = part.loc[2023], part.loc[2025]
        values = [subject, f"{b['mean']-a['mean']:+.3f}", f"{b['median']-a['median']:+.3f}",
                  f"{b['std']-a['std']:+.3f}", f"{b['skewness']-a['skewness']:+.3f}",
                  f"{b['kurtosis_excess']-a['kurtosis_excess']:+.3f}",
                  f"{b['pct_le_1']-a['pct_le_1']:+.2f}", f"{b['pct_lt_5']-a['pct_lt_5']:+.2f}",
                  f"{b['pct_ge_8']-a['pct_ge_8']:+.2f}", f"{b['pct_ge_9']-a['pct_ge_9']:+.2f}"]
        cells = table.add_row().cells
        for cell, value in zip(cells, values): cell.text = str(value)
    doc.add_heading("4. Biểu đồ và cách đọc", level=1)
    doc.add_paragraph("Histogram cho thấy tần suất tương đối theo các khoảng điểm; KDE làm mượt hình dạng và có thể che khuất các đỉnh nhỏ; boxplot so sánh trung vị và khoảng tứ phân vị; violin plot bổ sung hình dạng mật độ. Các biểu đồ dùng cùng thang điểm 0–10. Tiêu đề mô tả nội dung biểu đồ, không khẳng định trước kết quả.")
    for subject, paths in figures.items():
        doc.add_heading(subject, level=2)
        for path in paths:
            doc.add_picture(str(path), width=Inches(5.8))
            doc.add_paragraph(path.name)
    doc.add_heading("5. Nhận xét kết quả theo từng môn", level=1)
    doc.add_paragraph("Phần này được tạo tự động từ số liệu thực tế. Khi thuyết trình, hãy đọc nhận xét, chỉ vào bảng/biểu đồ tương ứng và nêu rõ đây là so sánh mô tả.")
    for subject in sorted(stats["subject"].unique()):
        doc.add_heading(subject, level=2)
        part = stats[stats["subject"] == subject].set_index("year")
        if 2023 not in part.index or 2025 not in part.index:
            doc.add_paragraph("Chỉ có dữ liệu hợp lệ của một trong hai năm; chưa đủ cơ sở so sánh trực tiếp.")
            continue
        a, b = part.loc[2023], part.loc[2025]
        dmean = b["mean"] - a["mean"]
        dmedian = b["median"] - a["median"]
        dstd = b["std"] - a["std"]
        direction = "tăng" if dmean > 0 else ("giảm" if dmean < 0 else "không đổi")
        med_direction = "tăng" if dmedian > 0 else ("giảm" if dmedian < 0 else "không đổi")
        spread = "phân tán hơn" if dstd > 0 else ("ít phân tán hơn" if dstd < 0 else "độ phân tán gần như không đổi")
        doc.add_paragraph(
            f"Số điểm hợp lệ: {int(a['n']):,} năm 2023 và {int(b['n']):,} năm 2025. "
            f"Điểm trung bình {direction} {abs(dmean):.3f} điểm (2023: {a['mean']:.3f}; 2025: {b['mean']:.3f}); "
            f"trung vị {med_direction} {abs(dmedian):.3f} điểm (2023: {a['median']:.3f}; 2025: {b['median']:.3f}). "
            f"Độ lệch chuẩn thay đổi {dstd:+.3f}, gợi ý phân phối {spread} theo chỉ số này."
        )
        doc.add_paragraph(
            f"Tỷ lệ điểm liệt (≤1) thay đổi {b['pct_le_1']-a['pct_le_1']:+.2f} điểm phần trăm; "
            f"dưới trung bình (<5) {b['pct_lt_5']-a['pct_lt_5']:+.2f}; "
            f"đạt từ 8 trở lên {b['pct_ge_8']-a['pct_ge_8']:+.2f}; "
            f"đạt từ 9 trở lên {b['pct_ge_9']-a['pct_ge_9']:+.2f} điểm phần trăm."
        )
        dskew = b["skewness"] - a["skewness"]
        dkurt = b["kurtosis_excess"] - a["kurtosis_excess"]
        center_note = ("cả trung bình và trung vị đều tăng" if dmean > 0 and dmedian > 0 else
                       "cả trung bình và trung vị đều giảm" if dmean < 0 and dmedian < 0 else
                       "trung bình và trung vị thay đổi không cùng chiều" if dmean * dmedian < 0 else
                       "trung tâm phân phối thay đổi không đồng nhất hoặc rất nhỏ")
        doc.add_paragraph(
            f"Về hình dạng, skewness thay đổi {dskew:+.3f} (2023: {a['skewness']:.3f}; 2025: {b['skewness']:.3f}); "
            f"excess kurtosis thay đổi {dkurt:+.3f} (2023: {a['kurtosis_excess']:.3f}; 2025: {b['kurtosis_excess']:.3f}). "
            "Skewness dương thường gợi ý đuôi phải dài hơn, âm gợi ý đuôi trái dài hơn; excess kurtosis dương/âm cần được đọc thận trọng "
            "như dấu hiệu về đuôi/độ nhọn, không dùng riêng để kết luận phân phối chuẩn."
        )
        doc.add_paragraph(
            f"Kết luận mô tả sơ bộ: {center_note}. "
            f"Chênh lệch tỷ lệ 2025–2023 lần lượt là ≤1: {b['pct_le_1']-a['pct_le_1']:+.2f}, "
            f"<5: {b['pct_lt_5']-a['pct_lt_5']:+.2f}, ≥8: {b['pct_ge_8']-a['pct_ge_8']:+.2f}, "
            f"≥9: {b['pct_ge_9']-a['pct_ge_9']:+.2f} điểm phần trăm. "
            "Để trả lời câu hỏi phổ điểm có dịch chuyển hay không, cần kiểm tra đồng thời histogram và KDE (vị trí các đỉnh/mật độ), "
            "boxplot (trung vị và IQR) và violin plot (hình dạng mật độ). Nếu các biểu đồ không cùng hướng, báo cáo phải nêu rõ sự khác biệt thay vì kết luận một chiều."
        )
        doc.add_paragraph(
            "Lời gợi ý khi trình bày: “Ở môn " + subject + ", khi so sánh 2023 với 2025, "
            f"điểm trung bình {direction}, còn trung vị {med_direction}; độ lệch chuẩn thay đổi {dstd:+.3f}. "
            f"{center_note.capitalize()}. "
            "Em đối chiếu histogram/KDE để xác định vùng điểm có mật độ dịch chuyển, rồi kiểm tra trung vị và IQR qua boxplot/violin. "
            "Các kết quả là khác biệt mô tả, chưa giải thích nguyên nhân.”"
        )
    doc.add_heading("6. Kịch bản trình bày gợi ý", level=1)
    doc.add_paragraph("Mở đầu: “Nhóm em phân tích mô tả và so sánh phân phối điểm thi THPT năm 2023 và 2025, dựa trên dữ liệu đã hợp mã tỉnh. Mục tiêu là xem vị trí, độ phân tán, hình dạng và tỷ lệ theo các ngưỡng điểm thay đổi như thế nào.”")
    doc.add_paragraph("Phương pháp: “Nhóm sử dụng mean, median, độ lệch chuẩn, skewness và excess kurtosis. Sau đó trực quan hóa bằng histogram, KDE, boxplot và violin plot. Các tỷ lệ được tính trên số điểm hợp lệ của từng môn-năm; không điền điểm thiếu và chỉ dùng điểm trong khoảng 0 đến 10.”")
    doc.add_paragraph("Khi trình bày kết quả: “Với mỗi môn, nhóm đọc bảng thống kê trước, sau đó kiểm tra các biểu đồ và đối chiếu tỷ lệ ≤1, <5, ≥8 và ≥9. Nhận xét cụ thể cho từng môn nằm trong mục 5 của báo cáo và được sinh từ dữ liệu thực tế khi chạy chương trình.”")
    doc.add_paragraph("Kết luận: “Phân tích cho biết phân phối điểm có khác biệt mô tả giữa hai năm hay không, nhưng không chứng minh nguyên nhân. Cần thận trọng vì đề thi, quy chế và thành phần thí sinh có thể khác nhau.”")
    doc.add_heading("7. Kết luận và giới hạn", level=1)
    doc.add_paragraph("Kết luận về dịch chuyển phổ điểm phải dựa đồng thời vào mean, median, độ phân tán, tỷ lệ ngưỡng và các biểu đồ. Không sử dụng riêng một chỉ số để đại diện toàn bộ phân phối.")
    doc.add_paragraph("Đây là phân tích mô tả, không chứng minh nguyên nhân của khác biệt và không nên diễn giải điểm số tăng/giảm như thay đổi trực tiếp về năng lực.")
    doc.save(output_path)

def run(file23, file25, output_dir):
    output_dir = Path(output_dir); figures_dir = output_dir / "figures"
    output_dir.mkdir(parents=True, exist_ok=True); figures_dir.mkdir(parents=True, exist_ok=True)
    data23 = tidy_year(read_dataset(file23), 2023)
    data25 = tidy_year(read_dataset(file25), 2025)
    subjects = sorted(set(data23.columns) & set(data25.columns) - {"year"})
    if not subjects:
        raise ValueError("Không có môn chung giữa hai dataset sau khi ánh xạ tên cột.")
    combined = pd.concat([data23[["year"] + subjects], data25[["year"] + subjects]], ignore_index=True)
    long = combined.melt(id_vars="year", var_name="subject", value_name="score").dropna(subset=["score"])
    stats = summarize(long)
    stats.to_csv(output_dir / "thong_ke_mo_ta_2023_2025.csv", index=False, encoding="utf-8-sig")
    long.to_csv(output_dir / "diem_long_2023_2025.csv", index=False, encoding="utf-8-sig")
    figures = {subject: plot_subject(long, subject, figures_dir) for subject in subjects}
    write_docx(stats, figures, output_dir / "bao_cao_phan_tich_pho_diem_2023_2025.docx",
               len(long), file23, file25)
    metadata = {"years": [2023, 2025], "subjects": subjects, "valid_score_rows": len(long),
                "source_2023": str(file23), "source_2025": str(file25),
                "note": "Output generated from local datasets; see Git LFS hydration status."}
    (output_dir / "run_metadata.json").write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Hoàn tất: {output_dir.resolve()} | {len(subjects)} môn | {len(long):,} điểm hợp lệ")

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--2023", dest="file23", default="data/processed/week3/exam_2023.parquet")
    parser.add_argument("--2025", dest="file25", default="data/dataset_diem_thi_thpt_2025.csv")
    parser.add_argument("--out", default="outputs/score_distribution")
    args = parser.parse_args()
    run(args.file23, args.file25, args.out)

if __name__ == "__main__":
    main()
